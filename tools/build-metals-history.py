#!/usr/bin/env python3
"""Generate market-history.json from the World Bank Pink Sheet monthly workbook.
The output contains measured monthly prices, never invented trading candles.
"""
from __future__ import annotations
import datetime as dt
import html
import json
import re
import urllib.request
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import load_workbook

HOME = "https://www.worldbank.org/en/research/commodity-markets"
FALLBACK = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
OUT = Path("data/market-history.json")
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; CebonkHistory/1.0; research)", "Accept": "*/*"}

def get(url: str) -> bytes:
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=35) as response:
        body = response.read(7_500_000)
        if len(body) >= 7_500_000:
            raise ValueError("World Bank workbook exceeds size limit")
        return body

def candidates() -> list[str]:
    urls = []
    try:
        page = get(HOME).decode("utf-8", errors="replace")
        for raw in re.findall(r'[^"\'\s<>]*CMO-Historical-Data-Monthly\.xlsx[^"\'\s<>]*', page, flags=re.I):
            url = urljoin(HOME, html.unescape(raw).replace("&amp;", "&"))
            if url.startswith("https://") and url not in urls:
                urls.append(url)
    except Exception as err:
        print("Discovery unavailable:", err)
    if FALLBACK not in urls:
        urls.append(FALLBACK)
    return urls

def parse_book(body: bytes) -> dict:
    if not body.startswith(b"PK"):
        raise ValueError("World Bank download is not an XLSX workbook")
    book = load_workbook(BytesIO(body), read_only=True, data_only=True)
    if "Monthly Prices" not in book.sheetnames:
        raise ValueError("World Bank workbook missing Monthly Prices sheet")
    sheet = book["Monthly Prices"]
    first = iter(sheet.values)
    sample = [next(first) for _ in range(12)]
    col = None
    for i, row in enumerate(sample):
        normalized = [re.sub(r"[^A-Z0-9_]", "", str(v or "").upper()) for v in row]
        gold = next((k for k, v in enumerate(normalized) if v == "GOLD"), None)
        silver = next((k for k, v in enumerate(normalized) if v == "SILVER"), None)
        if gold is not None and silver is not None:
            col = (gold, silver)
            break
    if col is None:
        raise ValueError("GOLD and SILVER data columns not found")
    series = {"XAUUSD": {}, "XAGUSD": {}}
    for row in list(sample) + list(first):
        date = row[0]
        if isinstance(date, (dt.datetime, dt.date)):
            month = date.strftime("%Y-%m")
        elif isinstance(date, str) and re.fullmatch(r"\d{4}M(0[1-9]|1[0-2])", date.strip()):
            month = date.strip().replace("M", "-")
        elif isinstance(date, str) and re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", date.strip()):
            month = date.strip()
        else:
            continue
        for sym, idx in [("XAUUSD", col[0]), ("XAGUSD", col[1])]:
            try:
                v = float(row[idx])
                if 0 < v < 1e7 and v == v:
                    series[sym][month] = round(v, 5)
            except (TypeError, ValueError, IndexError):
                pass
    data = {sym: [[m, v] for m, v in sorted(val.items())] for sym, val in series.items()}
    for sym, values in data.items():
        if len(values) < 200 or values[-1][0] < "2025-01":
            raise ValueError(f"Insufficient World Bank history for {sym}: {len(values)} rows")
    return data

def main():
    failures = []
    for url in candidates():
        try:
            body = get(url)
            data = parse_book(body)
            doc = {
                "version": 1,
                "source": "World Bank Pink Sheet — rata-rata bulanan USD/troy ounce",
                "source_url": url,
                "license": "CC BY 4.0",
                "retrieved": dt.datetime.now(dt.timezone.utc).date().isoformat(),
                "frequency": "monthly_average",
                "data": data,
            }
            OUT.parent.mkdir(parents=True, exist_ok=True)
            OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
            print("VALID", url, {k: (len(v), v[0][0], v[-1][0]) for k, v in data.items()})
            return
        except Exception as err:
            failures.append(f"{url}: {err}")
    raise SystemExit("World Bank history unavailable; refusing fabricated prices. " + " | ".join(failures))

if __name__ == "__main__":
    main()
